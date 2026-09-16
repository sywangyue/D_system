#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成三新展调研报告.xlsx (6 sheets)"""
import json, os, sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip3 install openpyxl")
    sys.exit(1)

BASE = os.path.dirname(os.path.abspath(__file__))

# Load data
results = [json.loads(l) for l in open(f"{BASE}/results.jsonl", encoding='utf-8')]
queue = json.load(open(f"{BASE}/queue.json", encoding='utf-8'))
roster_raw = json.load(open(f"{BASE}/roster.json", encoding='utf-8'))

# roster_raw is {attendee: [names], visitor: [names], exhibitor: [names]}
roster_sources = {}
src_labels = {'attendee': '参会', 'visitor': '观众', 'exhibitor': '展商'}
for src_key, names in roster_raw.items():
    label = src_labels.get(src_key, src_key)
    for name in names:
        if name not in roster_sources:
            roster_sources[name] = set()
        roster_sources[name].add(label)

roster_total = sum(len(v) for v in roster_raw.values())
print(f"Loaded: results={len(results)}, queue={len(queue)}, roster={roster_total}")

# Result map
result_map = {r['name']: r for r in results}

COLS = ['机构名称', '名单来源', '归属集团', '主体类型', '是否自办展', '自办展项目',
        '主赛道', '规模证据', 'UFI', '收购潜力', '战略契合', '合作切入点', '结论备注',
        'mwlab已知展会', '来源URL']

# === Sheet data ===
sheet1 = [r for r in results if r.get('ma', 0) >= 3 or r.get('fit', 0) >= 4]
sheet1.sort(key=lambda x: (-x.get('fit', 0), -x.get('ma', 0)))

sheet2 = [r for r in results if r.get('self') in ('是', '部分')]
sheet2.sort(key=lambda x: (-x.get('fit', 0), -x.get('ma', 0)))

sheet3 = [r for r in results if 1 <= r.get('ma', 0) <= 2 and r.get('fit', 0) >= 3]
sheet3.sort(key=lambda x: -x.get('fit', 0))

sheet4 = [r for r in results if r.get('self') == '待核']
def src_sort(r):
    s = roster_sources.get(r['name'], set())
    ss = '; '.join(s)
    if '参会' in ss: return 0
    if '展商' in ss: return 1
    return 2
sheet4.sort(key=src_sort)

excluded = [q for q in queue if q.get('需调研') != 'Y']

# Styles
header_font = Font(name='微软雅黑', bold=True, size=11, color='FFFFFF')
header_fill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
cell_align = Alignment(vertical='top', wrap_text=True)
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

ma_fills = {
    5: PatternFill(start_color='1B5E20', end_color='1B5E20', fill_type='solid'),
    4: PatternFill(start_color='2E7D32', end_color='2E7D32', fill_type='solid'),
    3: PatternFill(start_color='FDD835', end_color='FDD835', fill_type='solid'),
    2: PatternFill(start_color='EF6C00', end_color='EF6C00', fill_type='solid'),
    1: PatternFill(start_color='C62828', end_color='C62828', fill_type='solid'),
    0: PatternFill(start_color='BDBDBD', end_color='BDBDBD', fill_type='solid'),
}
ma_fonts = {5: Font(color='FFFFFF', bold=True, name='微软雅黑'),
            4: Font(color='FFFFFF', bold=True, name='微软雅黑'),
            3: Font(color='000000', name='微软雅黑'),
            2: Font(color='FFFFFF', name='微软雅黑'),
            1: Font(color='FFFFFF', name='微软雅黑'),
            0: Font(color='000000', name='微软雅黑')}

fit_fills = {
    5: PatternFill(start_color='0D47A1', end_color='0D47A1', fill_type='solid'),
    4: PatternFill(start_color='1565C0', end_color='1565C0', fill_type='solid'),
    3: PatternFill(start_color='FDD835', end_color='FDD835', fill_type='solid'),
    2: PatternFill(start_color='EF6C00', end_color='EF6C00', fill_type='solid'),
    1: PatternFill(start_color='BDBDBD', end_color='BDBDBD', fill_type='solid'),
    0: PatternFill(start_color='E0E0E0', end_color='E0E0E0', fill_type='solid'),
}
fit_fonts = {5: Font(color='FFFFFF', bold=True, name='微软雅黑'),
             4: Font(color='FFFFFF', bold=True, name='微软雅黑'),
             3: Font(color='000000', name='微软雅黑'),
             2: Font(color='FFFFFF', name='微软雅黑'),
             1: Font(color='000000', name='微软雅黑'),
             0: Font(color='000000', name='微软雅黑')}

FIELD_MAP = {
    '机构名称': 'name', '归属集团': 'group', '主体类型': 'type',
    '是否自办展': 'self', '自办展项目': 'projects', '主赛道': 'sectors',
    '规模证据': 'scale', '收购潜力': 'ma', '战略契合': 'fit',
    '来源URL': 'src', 'UFI': 'ufi'
}

def write_main_sheet(ws, data, title):
    ws.title = title
    for col_idx, col_name in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for row_idx, r in enumerate(data, 2):
        src = '; '.join(roster_sources.get(r['name'], ['']))
        for col_idx, col_name in enumerate(COLS, 1):
            if col_name == '名单来源':
                val = src
            elif col_name == '合作切入点':
                val = r.get('note', '')
            elif col_name == 'mwlab已知展会':
                val = ''
            elif col_name == '结论备注':
                val = ''
            else:
                key = FIELD_MAP.get(col_name, '')
                val = r.get(key, '') if key else ''

            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = cell_align
            cell.border = thin_border
            cell.font = Font(name='微软雅黑', size=10)

            if col_name == '收购潜力':
                try: v = int(val) if val != '' else 0
                except: v = 0
                cell.fill = ma_fills.get(v, ma_fills[0])
                cell.font = ma_fonts.get(v, ma_fonts[0])
                cell.alignment = Alignment(horizontal='center', vertical='top')
            elif col_name == '战略契合':
                try: v = int(val) if val != '' else 0
                except: v = 0
                cell.fill = fit_fills.get(v, fit_fills[0])
                cell.font = fit_fonts.get(v, fit_fonts[0])
                cell.alignment = Alignment(horizontal='center', vertical='top')

    ws.freeze_panes = 'A2'
    widths = {'A': 28, 'B': 12, 'C': 25, 'D': 14, 'E': 10, 'F': 50,
              'G': 18, 'H': 35, 'I': 8, 'J': 10, 'K': 10, 'L': 45,
              'M': 40, 'N': 30, 'O': 40}
    for cl, w in widths.items():
        ws.column_dimensions[cl].width = w
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(data)+1}"

# Create workbook
wb = Workbook()

ws1 = wb.active
write_main_sheet(ws1, sheet1, '1-重点标的')

ws2 = wb.create_sheet()
write_main_sheet(ws2, sheet2, '2-全部主办方')

ws3 = wb.create_sheet()
write_main_sheet(ws3, sheet3, '3-合作型机构')

ws4 = wb.create_sheet()
write_main_sheet(ws4, sheet4, '4-现场必问')

# Sheet 5: Excluded
ws5 = wb.create_sheet()
ws5.title = '5-已排除'
excl_cols = ['机构名称', '名单来源', '排除类别', '排除理由']
cat_map = {
    'tech_service': '科技/营销服务商',
    'execution': '执行类服务商',
    'hotel': '酒店/住宿文旅',
    'international': '国际同行（竞争对手）',
    'school': '院校'
}
for col_idx, cn in enumerate(excl_cols, 1):
    cell = ws5.cell(row=1, column=col_idx, value=cn)
    cell.font = header_font; cell.fill = header_fill
    cell.alignment = header_align; cell.border = thin_border

excluded_sorted = sorted(excluded, key=lambda x: x.get('排除类别', ''))
for row_idx, q in enumerate(excluded_sorted, 2):
    src = '; '.join(set(roster_sources.get(q['机构名称'], [''])))
    cat = cat_map.get(q.get('排除类别', ''), q.get('排除类别', ''))
    reason = q.get('排除理由', '')
    for col_idx, val in enumerate([q['机构名称'], src, cat, reason], 1):
        cell = ws5.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = cell_align; cell.border = thin_border
        cell.font = Font(name='微软雅黑', size=10)
ws5.freeze_panes = 'A2'
ws5.column_dimensions['A'].width = 30
ws5.column_dimensions['B'].width = 12
ws5.column_dimensions['C'].width = 22
ws5.column_dimensions['D'].width = 45

# Sheet 6: Full roster
ws6 = wb.create_sheet()
ws6.title = '6-原始名单'
orig_cols = ['机构名称', '名单来源', '是否调研', '调研结论']
for col_idx, cn in enumerate(orig_cols, 1):
    cell = ws6.cell(row=1, column=col_idx, value=cn)
    cell.font = header_font; cell.fill = header_fill
    cell.alignment = header_align; cell.border = thin_border

all_names = {}
for r in results:
    all_names[r['name']] = r.get('self', '')
for q in queue:
    if q['机构名称'] not in all_names:
        all_names[q['机构名称']] = '（已排除）'

sorted_names = sorted(all_names.items(), key=lambda x: x[0])
for row_idx, (name, conclusion) in enumerate(sorted_names, 2):
    src = '; '.join(set(roster_sources.get(name, [''])))
    researched = 'Y' if name in result_map else 'N（已排除）'
    for col_idx, val in enumerate([name, src, researched, conclusion], 1):
        cell = ws6.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = cell_align; cell.border = thin_border
        cell.font = Font(name='微软雅黑', size=10)
ws6.freeze_panes = 'A2'
ws6.column_dimensions['A'].width = 35
ws6.column_dimensions['B'].width = 12
ws6.column_dimensions['C'].width = 12
ws6.column_dimensions['D'].width = 15

output_path = f"{BASE}/三新展调研报告.xlsx"
wb.save(output_path)

print(f"\nSaved: {output_path}")
print(f"1-重点标的: {len(sheet1)}")
print(f"2-全部主办方: {len(sheet2)}")
print(f"3-合作型机构: {len(sheet3)}")
print(f"4-现场必问: {len(sheet4)}")
print(f"5-已排除: {len(excluded_sorted)}")
print(f"6-原始名单: {len(sorted_names)}")
